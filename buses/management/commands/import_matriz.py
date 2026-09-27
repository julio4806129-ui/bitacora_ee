"""
Importa MATRIZ.xlsx (hojas MATRIZ FLOTA + BUSES TMP) a InventarioFlota.

Uso:
  python manage.py import_matriz /ruta/MATRIZ.xlsx
  python manage.py import_matriz /ruta/MATRIZ.xlsx --solo-placas
  python manage.py import_matriz /ruta/MATRIZ.xlsx --solo-ee
"""
from datetime import datetime, date
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from buses.models import Patio, InventarioFlota, MovimientoFlota

# Mapeo COE del Excel → código Patio
COE_MAP = {
    'P-NDONA': 'LA_DONA',
    'P-DONA': 'LA_DONA',
    'P-PUEBLOS': 'LOS_PUEBLOS',
    'P-CHORRILLO': 'CHORRILLO',
    'P-OAGUA': 'OJO_DE_AGUA',
    'P-CABIMA': 'LA_CABIMA',
    'P-CURUNDU': 'CURUNDU',
}

# Nombres largos de MATRIZ FLOTA → código
PATIO_NOMBRE_MAP = {
    'PATIO LA DOÑA': 'LA_DONA',
    'PATIO LA DONA': 'LA_DONA',
    'PATIO LOS PUEBLOS': 'LOS_PUEBLOS',
    'PATIO CHORRILLO': 'CHORRILLO',
    'PATIO OJO DE AGUA': 'OJO_DE_AGUA',
    'PATIO LA CABIMA': 'LA_CABIMA',
    'PATIO CURUNDU': 'CURUNDU',
    'PATIO CURUNDÚ': 'CURUNDU',
}


def _s(val):
    if val is None:
        return ''
    return str(val).strip()


def _parse_date(val):
    if val is None or val == '':
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    text = str(val).strip()
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%Y/%m/%d'):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    return None


def _resolve_patio(coe_or_name, patio_cache):
    key = _s(coe_or_name).upper()
    cod = COE_MAP.get(key) or PATIO_NOMBRE_MAP.get(key)
    if not cod:
        # intentar match parcial
        for k, v in {**COE_MAP, **PATIO_NOMBRE_MAP}.items():
            if k in key or key in k:
                cod = v
                break
    if not cod:
        return None
    if cod not in patio_cache:
        patio_cache[cod] = Patio.objects.filter(codigo=cod).first()
    return patio_cache[cod]


class Command(BaseCommand):
    help = 'Importa MATRIZ.xlsx (flota + estado EE) a la base de datos'

    def add_arguments(self, parser):
        parser.add_argument('archivo', type=str, help='Ruta al archivo MATRIZ.xlsx')
        parser.add_argument('--solo-placas', action='store_true', help='Solo hoja MATRIZ FLOTA (patio/placa)')
        parser.add_argument('--solo-ee', action='store_true', help='Solo hoja BUSES TMP (estado EE)')
        parser.add_argument('--dry-run', action='store_true', help='Simular sin guardar')

    def handle(self, *args, **options):
        path = options['archivo']
        solo_placas = options['solo_placas']
        solo_ee = options['solo_ee']
        dry = options['dry_run']

        try:
            import openpyxl
        except ImportError:
            raise CommandError('Instale openpyxl: pip install openpyxl')

        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception as e:
            raise CommandError(f'No se pudo abrir {path}: {e}')

        self.stdout.write(f'Hojas: {wb.sheetnames}')
        patio_cache = {}
        stats = {'creados': 0, 'actualizados': 0, 'errores': 0}

        # Asegurar patios base
        for cod, nom, ord_ in [
            ('CHORRILLO', 'Chorrillo', 1),
            ('CURUNDU', 'Curundú', 2),
            ('LA_DONA', 'La Doña', 3),
            ('LOS_PUEBLOS', 'Los Pueblos', 4),
            ('LA_CABIMA', 'La Cabima', 5),
            ('OJO_DE_AGUA', 'Ojo de Agua', 6),
        ]:
            Patio.objects.update_or_create(codigo=cod, defaults={'nombre': nom, 'orden': ord_, 'activo': True})

        if not solo_ee and 'MATRIZ FLOTA' in wb.sheetnames:
            self.stdout.write('→ Importando MATRIZ FLOTA (placas)...')
            stats = self._import_placas(wb['MATRIZ FLOTA'], patio_cache, stats, dry)

        if not solo_placas and 'BUSES TMP' in wb.sheetnames:
            self.stdout.write('→ Importando BUSES TMP (estado EE)...')
            stats = self._import_ee(wb['BUSES TMP'], patio_cache, stats, dry)

        wb.close()
        self.stdout.write(self.style.SUCCESS(
            f"Listo. Creados={stats['creados']} Actualizados={stats['actualizados']} Errores={stats['errores']}"
            + (' [DRY-RUN]' if dry else '')
        ))

    def _import_placas(self, ws, patio_cache, stats, dry):
        rows = ws.iter_rows(min_row=2, values_only=True)  # header en fila 1 (0-index visual row2)
        # fila 0 del sheet es vacía a veces; header real en row index 1
        # openpyxl min_row=1 is first row. Header is row 2 in 1-based = index
        # Re-read with better detection
        all_rows = list(ws.iter_rows(values_only=True))
        header_idx = None
        for i, row in enumerate(all_rows[:5]):
            vals = [_s(c).upper() for c in (row or [])]
            if 'BUS' in vals and ('PLACA' in vals or 'COE' in vals):
                header_idx = i
                break
        if header_idx is None:
            self.stdout.write(self.style.WARNING('No se encontró header en MATRIZ FLOTA'))
            return stats

        headers = [_s(c).upper() for c in all_rows[header_idx]]
        col = {h: i for i, h in enumerate(headers) if h}

        for row in all_rows[header_idx + 1:]:
            if not row:
                continue
            try:
                bus_raw = row[col.get('BUS', 1)] if col.get('BUS') is not None else row[2]
                if bus_raw is None:
                    continue
                bus = int(float(str(bus_raw).strip()))
                if bus <= 0:
                    continue
                coe = _s(row[col['COE']]) if 'COE' in col else ''
                placa = _s(row[col['PLACA']]) if 'PLACA' in col else ''
                patio = _resolve_patio(coe, patio_cache)

                if dry:
                    stats['actualizados'] += 1
                    continue

                obj, created = InventarioFlota.objects.update_or_create(
                    bus_movil=bus,
                    defaults={
                        'placa': placa or '',
                        'coe': coe,
                        'patio_base': patio,
                        'patio_actual': patio.nombre if patio else '',
                    },
                )
                if created:
                    stats['creados'] += 1
                else:
                    stats['actualizados'] += 1
            except Exception as e:
                stats['errores'] += 1
                if stats['errores'] <= 5:
                    self.stdout.write(self.style.WARNING(f'  Error fila: {e}'))
        return stats

    def _import_ee(self, ws, patio_cache, stats, dry):
        all_rows = list(ws.iter_rows(values_only=True))
        header_idx = None
        for i, row in enumerate(all_rows[:6]):
            vals = [_s(c).upper() for c in (row or [])]
            if 'BUS' in vals and 'COE' in vals:
                header_idx = i
                break
        if header_idx is None:
            self.stdout.write(self.style.WARNING('No se encontró header en BUSES TMP'))
            return stats

        headers = [_s(c).upper() for c in all_rows[header_idx]]
        # columnas pueden repetirse (SISTEMA INSTALADO x2) — tomar índices por orden
        def find_col(name, occurrence=0):
            found = 0
            for i, h in enumerate(headers):
                if h == name:
                    if found == occurrence:
                        return i
                    found += 1
            return None

        c_coe = find_col('COE')
        c_bus = find_col('BUS')
        c_gen = find_col('ESTADO GENESIS')
        c_sis = find_col('ESTADO SISTEMA')
        c_sist1 = find_col('SISTEMA INSTALADO', 0)
        c_busae = find_col('ESTADO BUSAE')
        c_fw = find_col('ESTADO DE FW')
        c_boc = find_col('ESTADO BOCINA')
        c_adeb = find_col('ADECUACION BOCINA')
        c_adee = find_col('ADECUACION ELECTRICA')
        c_anc = find_col('ANCLAJE DE BOCINAS')
        c_sist2 = find_col('SISTEMA INSTALADO', 1)
        c_com = find_col('COMENTARIOS')
        c_mesi = find_col('MES INSTALADO')
        c_mesr = find_col('MES REVISADO')

        batch = []
        for row in all_rows[header_idx + 1:]:
            if not row or c_bus is None:
                continue
            try:
                bus_raw = row[c_bus]
                if bus_raw is None:
                    continue
                bus = int(float(str(bus_raw).strip()))
                if bus <= 0:
                    continue

                coe = _s(row[c_coe]) if c_coe is not None else ''
                patio = _resolve_patio(coe, patio_cache)

                defaults = {
                    'coe': coe,
                    'estado_genesis': _s(row[c_gen]) if c_gen is not None else '',
                    'estado_sistema': _s(row[c_sis]) if c_sis is not None else '',
                    'sistema_instalado': _s(row[c_sist1]) if c_sist1 is not None else '',
                    'estado_busae': _s(row[c_busae]) if c_busae is not None else '',
                    'estado_fw': _s(row[c_fw]) if c_fw is not None else '',
                    'estado_bocina': _s(row[c_boc]) if c_boc is not None else '',
                    'adecuacion_bocina': _s(row[c_adeb]) if c_adeb is not None else '',
                    'adecuacion_electrica': _s(row[c_adee]) if c_adee is not None else '',
                    'anclaje_bocinas': _s(row[c_anc]) if c_anc is not None else '',
                    'sistema_secundario': _s(row[c_sist2]) if c_sist2 is not None else '',
                    'comentarios': _s(row[c_com]) if c_com is not None else '',
                    'mes_instalado': _s(row[c_mesi]) if c_mesi is not None else '',
                    'mes_revisado': _s(row[c_mesr]) if c_mesr is not None else '',
                    'fecha_revision': _parse_date(row[c_mesr]) if c_mesr is not None else None,
                }
                if patio:
                    defaults['patio_base'] = patio
                    if not defaults.get('patio_actual'):
                        defaults['patio_actual'] = patio.nombre

                # Derivar estado_operativo
                tmp = InventarioFlota(bus_movil=bus, **{k: v for k, v in defaults.items() if k != 'patio_base'})
                tmp.estado_sistema = defaults['estado_sistema']
                tmp.estado_genesis = defaults['estado_genesis']
                defaults['estado_operativo'] = tmp.sincronizar_estado_operativo()

                if dry:
                    stats['actualizados'] += 1
                    continue

                obj, created = InventarioFlota.objects.update_or_create(
                    bus_movil=bus, defaults=defaults
                )
                if created:
                    stats['creados'] += 1
                else:
                    stats['actualizados'] += 1
            except Exception as e:
                stats['errores'] += 1
                if stats['errores'] <= 8:
                    self.stdout.write(self.style.WARNING(f'  Error bus row: {e}'))
        return stats
